import tempfile
import unittest
from pathlib import Path
from threading import Event

from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from converter.export import export_slides, validate_images
from converter.models import Options, ConversionError, Cancelled


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1] / 'tmp')
        self.root = Path(self.temp.name)
        self.images = []
        for i, color in enumerate(('red', 'green', 'blue')):
            path = self.root / f'{i}.png'
            im = Image.new('RGB', (800, 600), color)
            im.putpixel((10, 10), (255, 255, 255))
            im.save(path)
            self.images.append(path)

    def tearDown(self):
        self.temp.cleanup()

    def test_pdf_has_three_landscape_pages_in_order(self):
        out = self.root / 'тест.pdf'
        export_slides(self.images, out, Event())
        reader = PdfReader(out)
        self.assertEqual(len(reader.pages), 3)
        self.assertAlmostEqual(float(reader.pages[0].mediabox.width) / float(reader.pages[0].mediabox.height), 4 / 3)

    def test_pptx_preserves_image_order_and_ratio(self):
        out = self.root / 'тест.pptx'
        export_slides(self.images, out, Event())
        deck = Presentation(out)
        self.assertEqual(len(deck.slides), 3)
        self.assertAlmostEqual(deck.slide_width / deck.slide_height, 4 / 3)
        for slide, path in zip(deck.slides, self.images):
            self.assertEqual(len(slide.shapes), 1)
            self.assertEqual(slide.shapes[0].image.blob, path.read_bytes())

    def test_cancel_does_not_replace_existing_output(self):
        out = self.root / 'existing.pdf'
        out.write_bytes(b'original')
        stop = Event()
        stop.set()
        with self.assertRaises(Cancelled):
            export_slides(self.images, out, stop)
        self.assertEqual(out.read_bytes(), b'original')
        self.assertFalse(list(self.root.glob('*.part')))

    def test_blank_frame_rejected_but_identical_structural_slides_retained(self):
        blank = self.root / 'blank.png'
        Image.new('RGB', (200, 200), 'white').save(blank)
        with self.assertRaises(ConversionError):
            validate_images([blank])
        self.assertEqual(len(validate_images([self.images[0], self.images[0]])), 2)

    def test_options_reject_bad_manual_settings(self):
        source = self.root / 'slides.html'
        source.write_text('<html></html>')
        with self.assertRaises(ConversionError):
            Options(source, self.root / 'a.pdf', count=501).validate()
        with self.assertRaises(ConversionError):
            Options(source, source).validate()

    def test_forced_canvas_keeps_picture_ratio_and_uses_letterboxing(self):
        out = self.root / 'wide.pptx'
        export_slides(self.images, out, Event(), canvas_size=(1600, 900))
        deck = Presentation(out)
        self.assertAlmostEqual(deck.slide_width / deck.slide_height, 16 / 9, delta=0.000001)
        picture = deck.slides[0].shapes[0]
        self.assertAlmostEqual(picture.width / picture.height, 4 / 3, delta=0.000001)
        self.assertGreater(picture.left, 0)

    def test_pdf_encodes_pages_incrementally_and_can_cancel_between_them(self):
        out = self.root / 'incremental.pdf'
        out.write_bytes(b'previous result')
        stop = Event()
        checked = []

        def progress(current, total, status):
            partial = list(self.root.glob('.hpc-*.part'))
            self.assertEqual(len(partial), 1)
            self.assertEqual(len(PdfReader(partial[0]).pages), current)
            checked.append(current)
            if current == 2:
                stop.set()

        with self.assertRaises(Cancelled):
            export_slides(self.images, out, stop, progress)
        self.assertEqual(checked, [1, 2])
        self.assertEqual(out.read_bytes(), b'previous result')


if __name__ == '__main__':
    unittest.main()
