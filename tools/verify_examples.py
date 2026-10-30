"""Real exports + reopening + image/pixel checks. Run from project root."""
import json
import sys
from pathlib import Path
from threading import Event

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import pypdfium2 as pdfium
from PIL import Image, ImageChops, ImageStat
from pptx import Presentation
from pypdf import PdfReader
from converter.bridge import run_worker
from converter.export import export_slides
from converter.models import Options


def main():
    out = ROOT / 'test-results'
    out.mkdir(exist_ok=True)
    report = []
    thumbs = []
    examples = ('01-simple', '02-javascript', '03-animation', '04-resources', '05-ratio', '06-reveal')
    for name in examples:
        source = ROOT / 'examples' / (name + '.html')
        work = out / source.stem
        work.mkdir(exist_ok=True)
        config = work / 'config.json'
        options = Options(source, work / 'result.pdf', delay_ms=120)
        config.write_text(json.dumps(options.renderer_config(work), ensure_ascii=False), encoding='utf-8')
        capture = run_worker(config, lambda event: None, Event())
        frames = [Path(slide['file']) for slide in capture['slides']]
        expected = 4 if source.stem.startswith('06') else 3
        assert len(frames) == expected, source
        export_slides(frames, work / 'result.pdf', Event())
        export_slides(frames, work / 'result.pptx', Event())
        pdf = PdfReader(work / 'result.pdf')
        pptx = Presentation(work / 'result.pptx')
        assert len(pdf.pages) == expected
        assert len(pptx.slides) == expected
        doc = pdfium.PdfDocument(str(work / 'result.pdf'))
        pixel_errors = []
        for index, frame in enumerate(frames):
            # PPTX keeps the exact original PNG bytes, in order.
            assert pptx.slides[index].shapes[0].image.blob == frame.read_bytes()
            with Image.open(frame) as im:
                reference = im.convert('RGB')
            page = doc[index]
            assert abs(page.get_width() / page.get_height() - reference.width / reference.height) < 0.005
            assert abs(pptx.slide_width / pptx.slide_height - reference.width / reference.height) < 0.005
            bitmap = page.render(scale=reference.width / page.get_width())
            rendered = bitmap.to_pil().convert('RGB').resize(reference.size)
            # PDF is JPEG-based: compare pixel error, allowing compression noise.
            error = sum(ImageStat.Stat(ImageChops.difference(reference, rendered)).mean) / 3
            assert error < 12, (source, index, error)
            pixel_errors.append(round(error, 3))
            rendered.save(work / f'pdf-page-{index + 1}.png')
            thumb = rendered.copy()
            thumb.thumbnail((320, 240))
            thumbs.append((source.stem + f' / {index + 1}', thumb))
            rendered.close()
            reference.close()
            bitmap.close()
            page.close()
        doc.close()
        report.append({'example': source.name, 'count': expected, 'mode': capture['mode'],
                       'dimensions': [capture['width'], capture['height']], 'chrome': capture['chrome'],
                       'pdf_pages': len(pdf.pages), 'pptx_slides': len(pptx.slides),
                       'pptx_png_exact': True, 'pdf_mean_pixel_error': pixel_errors,
                       'text': [slide['text'] for slide in capture['slides']]})
        print(source.name, expected, 'PDF + PPTX verified', flush=True)
    from PIL import ImageDraw, ImageFont
    sheet = Image.new('RGB', (4 * 340, 5 * 280), '#edf1f5')
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 15)
    for i, (label, thumb) in enumerate(thumbs):
        x, y = (i % 4) * 340 + 10, (i // 4) * 280 + 10
        draw.text((x, y), label, fill='#203040', font=font)
        sheet.paste(thumb, (x, y + 25))
    sheet.save(out / 'contact-sheet.png')
    (out / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print('TOTAL:', sum(row['count'] for row in report), 'slides in each format')


if __name__ == '__main__':
    main()
