"""Raster exporters. Write atomically so cancellation never publishes partial output."""
import os
import tempfile
from pathlib import Path
from threading import Event

from PIL import Image, ImageOps, ImageStat
from pptx import Presentation
from pptx.util import Inches

from .models import Cancelled, ConversionError


def check_cancel(stop: Event):
    if stop.is_set():
        raise Cancelled('Конвертация отменена.')


def validate_images(paths, stop=None):
    if not paths:
        raise ConversionError('Браузер не вернул изображения слайдов.')
    result = []
    for item in paths:
        if stop is not None:
            check_cancel(stop)
        path = Path(item)
        try:
            with Image.open(path) as image:
                image.load()
                if min(image.size) < 20:
                    raise ConversionError(f'Слишком маленькое изображение: {path.name}.')
                # Only a truly uniform frame is rejected; flat backgrounds with text are valid.
                stat = ImageStat.Stat(image.convert('RGB'))
                if max(stat.stddev) < 0.05:
                    raise ConversionError(f'Слайд {len(result) + 1} пустой (одноцветный). Проверьте навигацию.')
        except (OSError, Image.DecompressionBombError) as exc:
            raise ConversionError(f'Повреждено изображение {path.name}: {exc}') from exc
        result.append(path)
    return result


def export_slides(paths, output: Path, stop: Event, progress=lambda *_: None, canvas_size=None):
    check_cancel(stop)
    paths = validate_images(paths, stop)
    output = Path(output)
    with Image.open(paths[0]) as first:
        size = canvas_size or first.size
    suffix = output.suffix.lower()
    if suffix not in ('.pdf', '.pptx'):
        raise ConversionError('Неизвестный формат экспорта.')
    fd, temporary = tempfile.mkstemp(prefix='.hpc-', suffix='.part', dir=output.parent)
    os.close(fd)
    try:
        if suffix == '.pdf':
            for i, path in enumerate(paths):
                check_cancel(stop)
                with Image.open(path) as source:
                    frame = source.convert('RGB')
                try:
                    if frame.size != size:
                        fitted = ImageOps.contain(frame, size, Image.Resampling.LANCZOS)
                        frame.close()
                        frame = fitted
                        canvas = Image.new('RGB', size, 'white')
                        canvas.paste(frame, ((size[0] - frame.width) // 2, (size[1] - frame.height) // 2))
                        frame.close()
                        frame = canvas
                    check_cancel(stop)
                    # Pillow appends one encoded page; never retain every decoded RGB slide.
                    frame.save(temporary, 'PDF', append=i > 0, resolution=144.0,
                               quality=92, subsampling=0, title=output.stem)
                finally:
                    frame.close()
                progress(i + 1, len(paths), 'Сборка PDF')
        else:
            deck = Presentation()
            deck.slide_width = Inches(13.333333)
            deck.slide_height = int(deck.slide_width * size[1] / size[0])
            for i, path in enumerate(paths):
                check_cancel(stop)
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                with Image.open(path) as image:
                    scale = min(deck.slide_width / image.width, deck.slide_height / image.height)
                    width, height = int(image.width * scale), int(image.height * scale)
                slide.shapes.add_picture(str(path), (deck.slide_width - width) // 2,
                                        (deck.slide_height - height) // 2, width=width, height=height)
                progress(i + 1, len(paths), 'Сборка PPTX')
            deck.save(temporary)
        check_cancel(stop)
        os.replace(temporary, output)
    except (OSError, ValueError) as exc:
        raise ConversionError(f'Не удалось сохранить результат: {exc}') from exc
    finally:
        Path(temporary).unlink(missing_ok=True)
    return output
